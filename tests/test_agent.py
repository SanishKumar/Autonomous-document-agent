import asyncio
import shutil
from pathlib import Path
from uuid import uuid4

import httpx
import pytest
from docx import Document
from docx.oxml.ns import qn

from app.agent import AgentService
from app.config import Settings
from app.llm import LLMHTTPError, extract_json_object
from app.main import create_app
from app.planner import Planner
from app.schemas import DocumentDraft, ExecutionPlan, ReviewResult
from app.tools import AgentState, AgentTools
from tests.fake_llm import FakeLLM


@pytest.fixture
def output_dir() -> Path:
    path = Path("outputs") / f".test-{uuid4().hex}"
    path.mkdir(parents=True)
    try:
        yield path
    finally:
        shutil.rmtree(path, ignore_errors=True)


def build_app(tmp_path: Path):
    settings = Settings(outputs_dir=tmp_path, auth_enabled=False)
    service = AgentService(settings, llm=FakeLLM())
    return create_app(service=service, settings=settings)


def call(app, method: str, url: str, **kwargs) -> httpx.Response:
    async def send() -> httpx.Response:
        transport = httpx.ASGITransport(app=app)
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
            return await client.request(method, url, **kwargs)

    return asyncio.run(send())


def assert_layout_tokens(document_path: Path) -> None:
    document = Document(document_path)
    section = document.sections[0]
    assert round(section.page_width.inches, 2) == 8.5
    assert round(section.page_height.inches, 2) == 11.0
    assert all(
        round(value.inches, 2) == 1.0
        for value in (
            section.top_margin,
            section.right_margin,
            section.bottom_margin,
            section.left_margin,
        )
    )
    assert document.styles["Normal"].font.name == "Calibri"
    assert document.styles["Normal"].font.size.pt == 11
    assert document.styles["Heading 1"].font.size.pt == 16

    for table in document.tables:
        table_width = table._tbl.tblPr.first_child_found_in("w:tblW")
        table_indent = table._tbl.tblPr.first_child_found_in("w:tblInd")
        assert table_width.get(qn("w:type")) == "dxa"
        assert int(table_width.get(qn("w:w"))) == 9360
        assert int(table_indent.get(qn("w:w"))) == 120
        grid_widths = [int(column.get(qn("w:w"))) for column in table._tbl.tblGrid.gridCol_lst]
        assert sum(grid_widths) == 9360
        header_marker = table.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader"))
        assert header_marker is not None
        for row in table.rows:
            assert row._tr.get_or_add_trPr().find(qn("w:trHeight")) is None
            cell_widths = [
                int(cell._tc.get_or_add_tcPr().get_or_add_tcW().get(qn("w:w")))
                for cell in row.cells
            ]
            assert cell_widths == grid_widths


def test_standard_request_generates_reviewed_docx(output_dir: Path) -> None:
    app = build_app(output_dir)
    response = call(
        app,
        "POST",
        "/agent",
        json={
            "request": (
                "Create a 90-day implementation plan for introducing an internal employee "
                "helpdesk chatbot in a company with 500 employees. Include objectives, project "
                "phases, owners, milestones, major risks, mitigations, and measurable success metrics."
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["status"] == "completed"
    assert payload["quality_score"] == 9
    assert payload["initial_quality_score"] == 9
    assert payload["revised"] is False
    assert len(payload["plan"]) == 5
    assert len(payload["execution_log"]) == 5
    assert all(task["status"] == "completed" for task in payload["execution_log"])

    document_path = output_dir / f"{payload['request_id']}.docx"
    assert document_path.is_file()
    assert_layout_tokens(document_path)
    document = Document(document_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "90-Day Implementation Plan" in text
    assert "Executive summary" in text
    assert "Success metrics" in text
    assert len(document.tables) >= 3

    download = call(app, "GET", payload["document_url"])
    assert download.status_code == 200
    assert download.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )


def test_complex_request_resolves_conflicts_and_revises(output_dir: Path) -> None:
    app = build_app(output_dir)
    response = call(
        app,
        "POST",
        "/agent",
        json={
            "request": (
                "Create a technical proposal for an autonomous customer support agent that must "
                "launch within 6 weeks. It must run completely on-premise, support Hindi and English, "
                "use only open-source models, remain within a budget of INR 3 lakh, support 10,000 "
                "concurrent users, and provide 99.99 percent availability. Important technical and "
                "organizational details are missing, so make reasonable assumptions and clearly "
                "identify conflicts, risks, and decisions."
            )
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["quality_score"] == 9
    assert payload["initial_quality_score"] == 7
    assert payload["revised"] is True
    assert len(payload["reflection_feedback"]) == 2
    assert len(payload["execution_log"]) == 6
    assert payload["execution_log"][1]["tool"] == "resolve_constraints"
    assert any("99.99" in constraint for constraint in payload["constraints"])
    assert any("pilot" in decision.casefold() for decision in payload["decisions"])

    document_path = output_dir / f"{payload['request_id']}.docx"
    assert_layout_tokens(document_path)
    document = Document(document_path)
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    table_text = "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    assert "Key decisions and trade-offs" in text
    assert "99.99%" in text
    assert "300,000" in table_text


def test_request_guardrail_rejects_non_meaningful_input(output_dir: Path) -> None:
    app = build_app(output_dir)
    response = call(
        app,
        "POST",
        "/agent",
        json={"request": "1234 5678 90"},
    )
    assert response.status_code == 422


def test_download_rejects_unsafe_filename(output_dir: Path) -> None:
    app = build_app(output_dir)
    response = call(app, "GET", "/documents/not-a-request.docx")
    assert response.status_code == 422


def test_json_normalizer_handles_markdown_fences() -> None:
    payload = extract_json_object('```json\n{"status": "ok", "count": 2}\n```')
    assert payload == {"status": "ok", "count": 2}


def test_groq_json_generation_400_is_retryable_but_other_400_is_not() -> None:
    generation_error = LLMHTTPError(
        400,
        '{"error":{"code":"json_validate_failed","message":"Failed to generate JSON"}}',
    )
    ordinary_error = LLMHTTPError(400, "invalid model parameter")

    assert generation_error.retryable is True
    assert ordinary_error.retryable is False


def test_planner_uses_bounded_fallback_after_two_invalid_plans() -> None:
    class InvalidPlannerLLM:
        calls = 0

        async def complete_json(self, system_prompt: str, user_prompt: str):
            self.calls += 1
            return {"title": "incomplete"}

    llm = InvalidPlannerLLM()
    result = asyncio.run(
        Planner(llm).create_plan(
            "Create a technical proposal with a strict budget and missing information."
        )
    )
    assert llm.calls == 2
    assert result.source == "fallback"
    assert result.attempts == 2
    assert result.plan.tasks[-1].tool == "generate_docx"
    assert any(task.tool == "resolve_constraints" for task in result.plan.tasks)


def test_planner_removes_unneeded_constraint_resolution_from_standard_plan() -> None:
    class OverplanningLLM:
        async def complete_json(self, system_prompt: str, user_prompt: str):
            return FakeLLM._plan(complex_request=True)

    result = asyncio.run(
        Planner(OverplanningLLM()).create_plan(
            "Create a 90-day implementation plan for an employee helpdesk chatbot."
        )
    )

    assert result.source == "model"
    assert len(result.plan.tasks) == 5
    assert [task.id for task in result.plan.tasks] == [
        "task_1",
        "task_2",
        "task_3",
        "task_4",
        "task_5",
    ]
    assert all(task.tool != "resolve_constraints" for task in result.plan.tasks)
    assert result.plan.tasks[-1].tool == "generate_docx"


def test_deterministic_review_rejects_vague_metrics_and_raw_identifiers() -> None:
    draft_payload = FakeLLM._draft(complex_request=False, revised=False)
    draft_payload["sections"][0]["paragraphs"] = [
        "pilot-project-for-500-employees",
        "John Doe as Technical_Owner confirmed a vendor SLA and a $10,000 upgrade "
        "based on 100 tickets per week.",
    ]
    metrics = next(
        section for section in draft_payload["sections"] if section["heading"] == "Success metrics"
    )
    metrics["paragraphs"] = []
    metrics["bullets"] = [
        "Reduce helpdesk tickets",
        "Increase employee satisfaction",
        "Improve chatbot adoption",
    ]
    metrics["table"] = None

    state = AgentState(
        request_id="test0001",
        original_request=(
            "Create an implementation plan with owners, milestones, risks, mitigations, "
            "and measurable success metrics."
        ),
        plan=ExecutionPlan.model_validate(FakeLLM._plan(complex_request=False)),
        draft=DocumentDraft.model_validate(draft_payload),
    )

    issues, instructions = AgentTools._deterministic_review_findings(state)
    assert any("raw slug" in issue for issue in issues)
    assert any("numeric" in issue for issue in issues)
    assert any("personal names" in issue for issue in issues)
    assert any("currency" in issue for issue in issues)
    assert any("ticket-volume" in issue for issue in issues)
    assert any("vendor" in issue for issue in issues)
    assert any("three numeric" in instruction for instruction in instructions)


def test_draft_normalizer_repairs_missing_conclusion() -> None:
    payload = FakeLLM._draft(complex_request=False, revised=False)
    payload.pop("conclusion")
    payload["executive_summary"] += " Reach <50%> adoption in <6> weeks."
    payload["sections"][0]["bullets"].append("Achieve 80% accuracy as an objective")
    payload["sections"][0]["bullets"].append("Document Revisor owns plan maintenance")
    metrics = next(
        section for section in payload["sections"] if section["heading"] == "Success metrics"
    )
    metrics["table"]["rows"][0][1] = 2

    repaired = AgentTools._normalize_payload(DocumentDraft, payload)
    draft = DocumentDraft.model_validate(repaired)

    assert "decision gates" in draft.conclusion
    assert "Reach 50% adoption in 6 weeks" in draft.executive_summary
    assert all("80% accuracy" not in item for item in draft.sections[0].bullets)
    assert "Project Lead owns plan maintenance" in draft.sections[0].bullets
    metrics_section = next(
        section for section in draft.sections if section.heading == "Success metrics"
    )
    assert metrics_section.table.rows[0][1] == "2"


def test_review_filter_keeps_requested_defects_and_drops_unrequested_choices() -> None:
    review = ReviewResult(
        score=7,
        passed=False,
        missing_items=[
            "Select a specific chatbot platform.",
            "Success metrics need numeric targets.",
        ],
        contradictions=[],
        revision_instructions=[
            "Add a resource allocation plan for limited IT resources.",
            "Add measurable targets and measurement methods.",
        ],
    )

    filtered = AgentTools._remove_non_blocking_unknowns(
        review,
        "Create an implementation plan with measurable success metrics.",
    )

    assert filtered.missing_items == ["Success metrics need numeric targets."]
    assert filtered.revision_instructions == ["Add measurable targets and measurement methods."]
