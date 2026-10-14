from __future__ import annotations

import math
from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from app.schemas import ConstraintResolution, DocumentDraft, DocumentTable, ExecutionPlan

# Resolved standard_business_brief tokens from the document design specification.
PAGE_WIDTH_IN = 8.5
PAGE_HEIGHT_IN = 11.0
MARGIN_IN = 1.0
HEADER_FOOTER_IN = 0.492
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS_DXA = {"top": 80, "bottom": 80, "start": 120, "end": 120}

FONT = "Calibri"
INK = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "667085"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"


class DocumentBuilder:
    """Transforms structured content into a deterministic, styled Word document."""

    def __init__(self, output_dir: Path) -> None:
        self.output_dir = output_dir

    def build(
        self,
        *,
        request_id: str,
        plan: ExecutionPlan,
        draft: DocumentDraft,
        assumptions: list[str],
        resolution: ConstraintResolution | None,
    ) -> Path:
        self.output_dir.mkdir(parents=True, exist_ok=True)
        output_path = self.output_dir / f"{request_id}.docx"

        document = Document()
        self._configure_document(document)
        bullet_num_id = self._create_numbering(document, kind="bullet")
        decimal_num_id = self._create_numbering(document, kind="decimal")
        self._add_header_footer(document, plan.document_type, request_id)
        self._add_title_block(document, plan, request_id)

        document.add_heading("Executive summary", level=1)
        self._add_callout(document, draft.executive_summary)

        if assumptions:
            document.add_heading("Assumptions", level=1)
            intro = document.add_paragraph(
                "The following assumptions make the recommendation actionable and should be "
                "validated during discovery."
            )
            intro.paragraph_format.keep_with_next = True
            for item in assumptions:
                self._add_list_item(document, item, bullet_num_id)

        if resolution is not None:
            document.add_heading("Key decisions and trade-offs", level=1)
            document.add_paragraph(resolution.recommended_scope)
            decision_table = DocumentTable(
                headers=["Issue", "Decision", "Impact"],
                rows=[
                    [decision.issue, decision.decision, decision.impact]
                    for decision in resolution.decisions
                ],
            )
            self._add_caption(document, "Decision record")
            self._add_table(document, decision_table)
            if resolution.deferred_items:
                label = document.add_paragraph()
                label.paragraph_format.space_before = Pt(8)
                label.paragraph_format.space_after = Pt(4)
                run = label.add_run("Deferred beyond the recommended scope")
                run.bold = True
                for item in resolution.deferred_items:
                    self._add_list_item(document, item, bullet_num_id)

        for section in draft.sections:
            document.add_heading(section.heading, level=1)
            for text in section.paragraphs:
                document.add_paragraph(text)
            for item in section.bullets:
                self._add_list_item(document, item, bullet_num_id)
            for item in section.numbered_items:
                self._add_list_item(document, item, decimal_num_id)
            if section.table is not None:
                self._add_caption(document, section.heading)
                self._add_table(document, section.table)

        has_closing_section = any(
            "conclusion" in section.heading.casefold()
            or "recommendation" in section.heading.casefold()
            for section in draft.sections
        )
        if not has_closing_section:
            document.add_heading("Conclusion", level=1)
            document.add_paragraph(draft.conclusion)

        document.core_properties.title = plan.title
        document.core_properties.subject = plan.document_type
        document.core_properties.author = "Autonomous Business Document Agent"
        document.core_properties.comments = (
            "Generated from an explicit request, validated plan, bounded tools, and quality review."
        )
        document.core_properties.created = datetime.now(UTC)
        document.core_properties.modified = datetime.now(UTC)
        temporary = output_path.with_suffix(f".{uuid4().hex}.tmp")
        try:
            document.save(temporary)
            temporary.replace(output_path)
        finally:
            temporary.unlink(missing_ok=True)
        return output_path

    def _configure_document(self, document: DocumentObject) -> None:
        section = document.sections[0]
        section.page_width = Inches(PAGE_WIDTH_IN)
        section.page_height = Inches(PAGE_HEIGHT_IN)
        section.top_margin = Inches(MARGIN_IN)
        section.right_margin = Inches(MARGIN_IN)
        section.bottom_margin = Inches(MARGIN_IN)
        section.left_margin = Inches(MARGIN_IN)
        section.header_distance = Inches(HEADER_FOOTER_IN)
        section.footer_distance = Inches(HEADER_FOOTER_IN)

        styles = document.styles
        normal = styles["Normal"]
        self._set_style_font(normal, FONT, 11, "000000")
        normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        normal.paragraph_format.space_before = Pt(0)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.10

        title = styles["Title"]
        self._set_style_font(title, FONT, 26, INK, bold=True)
        title.paragraph_format.space_before = Pt(0)
        title.paragraph_format.space_after = Pt(6)
        title.paragraph_format.line_spacing = 1.0
        title.paragraph_format.keep_with_next = True

        subtitle = styles["Subtitle"]
        self._set_style_font(subtitle, FONT, 10.5, MUTED)
        subtitle.paragraph_format.space_before = Pt(0)
        subtitle.paragraph_format.space_after = Pt(12)
        subtitle.paragraph_format.keep_with_next = True

        heading_tokens = {
            "Heading 1": (16, BLUE, 16, 8),
            "Heading 2": (13, BLUE, 12, 6),
            "Heading 3": (12, DARK_BLUE, 8, 4),
        }
        for name, (size, color, before, after) in heading_tokens.items():
            style = styles[name]
            self._set_style_font(style, FONT, size, color, bold=True)
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.keep_together = True

        for name in ("List Bullet", "List Number"):
            style = styles[name]
            self._set_style_font(style, FONT, 11, "000000")
            style.paragraph_format.left_indent = Inches(0.5)
            style.paragraph_format.first_line_indent = Inches(-0.25)
            style.paragraph_format.space_after = Pt(8)
            style.paragraph_format.line_spacing = 1.167

        caption = styles["Caption"]
        self._set_style_font(caption, FONT, 9, MUTED, bold=True)
        caption.paragraph_format.space_before = Pt(8)
        caption.paragraph_format.space_after = Pt(4)
        caption.paragraph_format.keep_with_next = True

    @staticmethod
    def _set_style_font(
        style: object,
        name: str,
        size: float,
        color: str,
        *,
        bold: bool | None = None,
    ) -> None:
        style.font.name = name
        style._element.rPr.rFonts.set(qn("w:ascii"), name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), name)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        if bold is not None:
            style.font.bold = bold

    def _add_header_footer(
        self, document: DocumentObject, document_type: str, request_id: str
    ) -> None:
        section = document.sections[0]
        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
        left = paragraph.add_run(document_type.upper())
        self._format_run(left, 8, MUTED, bold=True)
        right = paragraph.add_run("\tBUSINESS DOCUMENT")
        self._format_run(right, 8, MUTED)

        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
        left = paragraph.add_run(f"Request {request_id}")
        self._format_run(left, 8, MUTED)
        right = paragraph.add_run("\tPage ")
        self._format_run(right, 8, MUTED)
        self._add_page_field(paragraph)

    def _add_title_block(
        self, document: DocumentObject, plan: ExecutionPlan, request_id: str
    ) -> None:
        kicker = document.add_paragraph()
        kicker.paragraph_format.space_before = Pt(12)
        kicker.paragraph_format.space_after = Pt(3)
        kicker.paragraph_format.keep_with_next = True
        run = kicker.add_run(plan.document_type.upper())
        self._format_run(run, 9, BLUE, bold=True)

        document.add_paragraph(plan.title, style="Title")
        prepared = datetime.now().astimezone().strftime("%d %B %Y")
        subtitle = document.add_paragraph(
            f"Prepared for {plan.audience}  |  {prepared}  |  Request {request_id}",
            style="Subtitle",
        )
        subtitle.paragraph_format.keep_with_next = True
        rule = document.add_paragraph()
        rule.paragraph_format.space_before = Pt(0)
        rule.paragraph_format.space_after = Pt(10)
        rule.paragraph_format.keep_with_next = True
        self._set_paragraph_bottom_border(rule, BLUE, size=12)

    def _add_callout(self, document: DocumentObject, text: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.12)
        paragraph.paragraph_format.right_indent = Inches(0.08)
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(10)
        paragraph.paragraph_format.line_spacing = 1.10
        self._shade_paragraph(paragraph, CALLOUT)
        self._set_paragraph_left_border(paragraph, BLUE, size=18)
        run = paragraph.add_run(text)
        self._format_run(run, 11, INK)

    def _add_list_item(self, document: DocumentObject, text: str, num_id: int) -> None:
        style_name = "List Bullet" if self._is_bullet_num(document, num_id) else "List Number"
        paragraph = document.add_paragraph(style=style_name)
        paragraph.add_run(text)
        p_pr = paragraph._p.get_or_add_pPr()
        num_pr = p_pr.get_or_add_numPr()
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_element = OxmlElement("w:numId")
        num_id_element.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(num_id_element)

    @staticmethod
    def _is_bullet_num(document: DocumentObject, num_id: int) -> bool:
        return getattr(document, "_agent_bullet_num_id", None) == num_id

    def _create_numbering(self, document: DocumentObject, *, kind: str) -> int:
        numbering = document.part.numbering_part.element
        abstract_ids = [
            int(item.get(qn("w:abstractNumId"))) for item in numbering.findall(qn("w:abstractNum"))
        ]
        num_ids = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
        abstract_id = max(abstract_ids, default=0) + 1
        num_id = max(num_ids, default=0) + 1

        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)
        number_format = OxmlElement("w:numFmt")
        number_format.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
        level.append(number_format)
        level_text = OxmlElement("w:lvlText")
        level_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
        level.append(level_text)
        justification = OxmlElement("w:lvlJc")
        justification.set(qn("w:val"), "left")
        level.append(justification)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        p_pr.append(tabs)
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), "720")
        indent.set(qn("w:hanging"), "360")
        p_pr.append(indent)
        level.append(p_pr)
        run_properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), FONT)
        fonts.set(qn("w:hAnsi"), FONT)
        run_properties.append(fonts)
        level.append(run_properties)
        abstract.append(level)
        numbering.append(abstract)

        number = OxmlElement("w:num")
        number.set(qn("w:numId"), str(num_id))
        abstract_reference = OxmlElement("w:abstractNumId")
        abstract_reference.set(qn("w:val"), str(abstract_id))
        number.append(abstract_reference)
        numbering.append(number)
        if kind == "bullet":
            document._agent_bullet_num_id = num_id
        return num_id

    def _add_caption(self, document: DocumentObject, label: str) -> None:
        document.add_paragraph(f"Table: {label}", style="Caption")

    def _add_table(self, document: DocumentObject, data: DocumentTable) -> None:
        table = document.add_table(rows=1, cols=len(data.headers))
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        table.style = "Table Grid"
        for index, header in enumerate(data.headers):
            table.rows[0].cells[index].text = header
        for source_row in data.rows:
            cells = table.add_row().cells
            for index, value in enumerate(source_row):
                cells[index].text = value

        widths = self._calculate_column_widths(data)
        self._set_table_geometry(table, widths)
        self._set_repeat_table_header(table.rows[0])
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                self._set_cell_margins(cell)
                if row_index == 0:
                    self._shade_cell(cell, LIGHT_GRAY)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.05
                    if self._short_value_column(data, cell_index):
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        self._format_run(
                            run,
                            9 if len(data.headers) >= 5 else 9.5,
                            INK if row_index == 0 else "000000",
                            bold=row_index == 0,
                        )
        spacer = document.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)

    @staticmethod
    def _short_value_column(data: DocumentTable, index: int) -> bool:
        values = [row[index] for row in data.rows]
        header = data.headers[index].casefold()
        short_header = any(
            marker in header
            for marker in ("date", "week", "owner", "status", "score", "phase", "duration")
        )
        return short_header and max((len(value) for value in values), default=0) <= 28

    @staticmethod
    def _calculate_column_widths(data: DocumentTable) -> list[int]:
        column_count = len(data.headers)
        if column_count == 1:
            return [CONTENT_WIDTH_DXA]
        if column_count == 2:
            lengths = [
                max(len(data.headers[index]), *(len(row[index]) for row in data.rows))
                for index in range(2)
            ]
            if lengths[0] * 1.8 < lengths[1]:
                return [2700, 6660]

        max_lengths = [
            max(len(data.headers[index]), *(len(row[index]) for row in data.rows))
            for index in range(column_count)
        ]
        weights = [min(3.0, max(0.9, math.sqrt(length) / 3.2)) for length in max_lengths]
        minimum = 1080 if column_count >= 5 else 1250
        available = CONTENT_WIDTH_DXA - (minimum * column_count)
        weight_total = sum(weights)
        widths = [minimum + round(available * weight / weight_total) for weight in weights]
        widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
        return widths

    @staticmethod
    def _set_table_geometry(table: Table, widths: list[int]) -> None:
        table_pr = table._tbl.tblPr
        layout = table_pr.first_child_found_in("w:tblLayout")
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            table_pr.append(layout)
        layout.set(qn("w:type"), "fixed")
        table_width = table_pr.first_child_found_in("w:tblW")
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
        table_indent = table_pr.first_child_found_in("w:tblInd")
        if table_indent is None:
            table_indent = OxmlElement("w:tblInd")
            table_pr.append(table_indent)
        table_indent.set(qn("w:type"), "dxa")
        table_indent.set(qn("w:w"), str(TABLE_INDENT_DXA))

        grid = table._tbl.tblGrid
        for child in list(grid):
            grid.remove(child)
        for width in widths:
            column = OxmlElement("w:gridCol")
            column.set(qn("w:w"), str(width))
            grid.append(column)
        for row in table.rows:
            for index, cell in enumerate(row.cells):
                cell.width = Inches(widths[index] / 1440)
                tc_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
                tc_width.set(qn("w:type"), "dxa")
                tc_width.set(qn("w:w"), str(widths[index]))

    @staticmethod
    def _set_cell_margins(cell: _Cell) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        margins = tc_pr.first_child_found_in("w:tcMar")
        if margins is None:
            margins = OxmlElement("w:tcMar")
            tc_pr.append(margins)
        for edge, value in CELL_MARGINS_DXA.items():
            element = margins.find(qn(f"w:{edge}"))
            if element is None:
                element = OxmlElement(f"w:{edge}")
                margins.append(element)
            element.set(qn("w:w"), str(value))
            element.set(qn("w:type"), "dxa")

    @staticmethod
    def _set_repeat_table_header(row: object) -> None:
        tr_pr = row._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        tr_pr.append(repeat)

    @staticmethod
    def _shade_cell(cell: _Cell, color: str) -> None:
        tc_pr = cell._tc.get_or_add_tcPr()
        shading = tc_pr.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            tc_pr.append(shading)
        shading.set(qn("w:fill"), color)

    @staticmethod
    def _shade_paragraph(paragraph: Paragraph, color: str) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), color)
        p_pr.append(shading)

    @staticmethod
    def _set_paragraph_bottom_border(paragraph: Paragraph, color: str, *, size: int) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), str(size))
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), color)
        borders.append(bottom)
        p_pr.append(borders)

    @staticmethod
    def _set_paragraph_left_border(paragraph: Paragraph, color: str, *, size: int) -> None:
        p_pr = paragraph._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), str(size))
        left.set(qn("w:space"), "6")
        left.set(qn("w:color"), color)
        borders.append(left)
        p_pr.append(borders)

    @staticmethod
    def _add_page_field(paragraph: Paragraph) -> None:
        run = paragraph.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instruction = OxmlElement("w:instrText")
        instruction.set(qn("xml:space"), "preserve")
        instruction.text = " PAGE "
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        text = OxmlElement("w:t")
        text.text = "1"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run._r.extend([begin, instruction, separate, text, end])
        run.font.name = FONT
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)

    @staticmethod
    def _format_run(
        run: object,
        size: float,
        color: str,
        *,
        bold: bool | None = None,
    ) -> None:
        run.font.name = FONT
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT)
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)
        if bold is not None:
            run.bold = bold
